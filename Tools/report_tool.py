from pathlib import Path
from docx import Document
from langchain.tools import tool
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet



OUTPUT_DIR = Path("outputs")
OUTPUT_DIR.mkdir(exist_ok=True)

# Export tool in word Doc
@tool
def create_docx_report(title: str, content: str, file_name: str):
    """Create a Word report from a title and content."""

    document = Document()
    document.add_heading(title, level=1)

    for paragraph in content.split("\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph)
    file_path = OUTPUT_DIR / file_name

    if not str(file_path).endswith(".docx"):
        file_path = file_path.with_suffix(".docx")
    document.save(file_path)

    return f"Report created successfully: {file_path}"

# Export tool in PDF
@tool
def create_pdf_report(title: str, content: str, file_name: str):
    """Create a PDF report."""

    if not file_name.endswith(".pdf"):
        file_name += ".pdf"

    file_path = OUTPUT_DIR / file_name
    styles = getSampleStyleSheet()

    elements = [
        Paragraph(title, styles["Title"]),
        Spacer(1, 20)
    ]

    for paragraph in content.split("\n"):
        if paragraph.strip():
            elements.append(
                Paragraph(paragraph, styles["BodyText"])
            )
            elements.append(Spacer(1, 10))

    document = SimpleDocTemplate(
        str(file_path),
        pagesize=letter
    )

    document.build(elements)

    return f"PDF created successfully: {file_path}"